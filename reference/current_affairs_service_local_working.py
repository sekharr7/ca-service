import asyncio
import base64
import csv
import json
import os
import random
import re
import shutil
import sys
import time
import uuid
from datetime import date, datetime
from io import BytesIO
from typing import List, Tuple
from urllib.parse import urlparse

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(line_buffering=True)


# Ensure crawl4ai runtime site-packages are loaded
CRAWL4AI_RUNTIME = "/Volumes/WD_1TB/LLM/runtimes/crawl4ai/.venv/lib/python3.13/site-packages"
if os.path.exists(CRAWL4AI_RUNTIME) and CRAWL4AI_RUNTIME not in sys.path:
    sys.path.insert(0, CRAWL4AI_RUNTIME)

# Playwright & Chromium local configuration
os.environ.setdefault(
    "PLAYWRIGHT_BROWSERS_PATH",
    "/Volumes/WD_1TB/LLM/cache/playwright"
)

CHROMIUM_EXECUTABLE_PATH = (
    "/Volumes/WD_1TB/LLM/cache/playwright/chromium-1234/"
    "chrome-mac-arm64/Google Chrome for Testing.app/"
    "Contents/MacOS/Google Chrome for Testing"
)

if os.path.exists(CHROMIUM_EXECUTABLE_PATH):
    os.environ["CHROME_BIN"] = CHROMIUM_EXECUTABLE_PATH

from dotenv import load_dotenv
load_dotenv()

from azure.storage.blob import BlobServiceClient
from bs4 import BeautifulSoup
from openai import AzureOpenAI, OpenAI
from PIL import Image as PILImage
from sqlalchemy import Column, DateTime, Integer, String, Text, create_engine, orm
from sqlalchemy.dialects.mysql import LONGTEXT, TINYINT
from sqlalchemy.orm import declarative_base, sessionmaker

from crawl4ai import AsyncWebCrawler, BrowserConfig, CacheMode, CrawlerRunConfig
from crawl4ai.content_filter_strategy import BM25ContentFilter
from crawl4ai.markdown_generation_strategy import DefaultMarkdownGenerator

# ── Local Models Definition (mapped to dev database) ───────────────────
Base = declarative_base()

class CurrentAffair(Base):
    __tablename__ = "mtr_current_affairs_edk"
    id = Column(Integer, primary_key=True, autoincrement=True)
    topic = Column(String(255), nullable=False)
    content = Column(LONGTEXT, nullable=False)
    theme = Column(Text, nullable=True)
    theme_topic = Column(Text, nullable=True)
    theme_subtopic = Column(Text, nullable=True)
    ca_created_date = Column(DateTime, nullable=False)
    created_at = Column(DateTime, nullable=True)
    source = Column(Text, nullable=True)
    view_count = Column(Integer, nullable=True)
    pinecone_synced = Column(TINYINT, nullable=False, default=0)

class Highlight(Base):
    __tablename__ = "mtr_ca_highlights_edk"
    id = Column(Integer, primary_key=True)
    current_affair_id = Column(Integer, nullable=False)
    highlight_text = Column(LONGTEXT, nullable=False)
    created_at = Column(DateTime, nullable=True)

class Image(Base):
    __tablename__ = "mtr_ca_images_edk"
    id = Column(Integer, primary_key=True)
    current_affair_id = Column(Integer, nullable=False)
    image_url = Column(Text, nullable=False)
    created_at = Column(DateTime, nullable=True)

class Keyword(Base):
    __tablename__ = "mtr_ca_kword_edk"
    id = Column(Integer, primary_key=True, autoincrement=True)
    current_affair_id = Column(Integer, nullable=False)
    kword = Column(Text, nullable=False)

class BulkJob(Base):
    __tablename__ = "mtr_ca_bulk_jobs"
    id = Column(String(36), primary_key=True)
    status = Column(String(20), nullable=False, default="pending")
    total = Column(Integer, nullable=False, default=0)
    processed = Column(Integer, nullable=False, default=0)
    succeeded = Column(Integer, nullable=False, default=0)
    failed = Column(Integer, nullable=False, default=0)
    skipped = Column(Integer, nullable=False, default=0)
    details = Column(LONGTEXT, nullable=True)
    started_at = Column(DateTime, nullable=True)
    completed_at = Column(DateTime, nullable=True)
    created_at = Column(DateTime, nullable=True)

# ── Current Affairs Generation – env config from .env ──────────────────
DB_URI = os.getenv("DB_URI")
CA_GEN_AZURE_OPENAI_ENDPOINT = os.getenv("CA_GEN_AZURE_OPENAI_ENDPOINT")
CA_GEN_AZURE_OPENAI_API_KEY = os.getenv("CA_GEN_AZURE_OPENAI_API_KEY")
CA_GEN_AZURE_OPENAI_API_VERSION = os.getenv("CA_GEN_AZURE_OPENAI_API_VERSION", "2025-04-01-preview")
CA_GEN_AZURE_OPENAI_IMAGE_DEPLOYMENT = os.getenv("CA_GEN_AZURE_OPENAI_IMAGE_DEPLOYMENT", "gpt-image-2")
AZURE_OPENAI_IMAGE_DEPLOYMENT = CA_GEN_AZURE_OPENAI_IMAGE_DEPLOYMENT
PPLEX_API_KEY = os.getenv("CA_GEN_PPLEX_API_KEY") or os.getenv("PPLEX_API_KEY")
AZ_CONN_STR = os.getenv("CA_GEN_AZURE_STORAGE_CONNECTION_STRING") or os.getenv("AZURE_STORAGE_CONNECTION_STRING")
AZ_CONTAINER = os.getenv("CA_GEN_AZURE_CONTAINER_NAME", "thumbnail")

IMAGE_OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "generated_images")
os.makedirs(IMAGE_OUTPUT_DIR, exist_ok=True)

image_client = AzureOpenAI(
    api_key=CA_GEN_AZURE_OPENAI_API_KEY,
    api_version=CA_GEN_AZURE_OPENAI_API_VERSION,
    azure_endpoint=CA_GEN_AZURE_OPENAI_ENDPOINT,
    timeout=90.0,
    max_retries=2,
)

blob_service_client = BlobServiceClient.from_connection_string(AZ_CONN_STR)

_engine = None
_SessionFactory = None

def get_db_session() -> orm.Session:
    global _engine, _SessionFactory
    if _engine is None:
        if not DB_URI:
            raise ValueError("DB_URI is not set in .env")
        _engine = create_engine(DB_URI, pool_recycle=3600, pool_pre_ping=True)
        _SessionFactory = sessionmaker(autocommit=False, autoflush=False, bind=_engine)
    return _SessionFactory()

def gen_image_prompt(topic: str, summary: str) -> str:
    return f"""
{topic}

{summary[:500]}

You're an AI capable of converting textual information into concrete images and your task is to convert the information in this article into symbolic images as you understand them.
Create a clean educational infographic for UPSC aspirants.

Style guidelines:
- Modern newspaper explainer style
- Structured infographic layout with visual sections
- Educational and fact-focused presentation
- Use icons, symbols, charts, policy visuals, governance visuals, science visuals, or institutional imagery when relevant
- Clear typography and concise educational labels are allowed
- Flat infographic or polished editorial infographic style
- Avoid cinematic photorealism
- Avoid fictional or highly detailed geographic maps
- Prefer schematic or symbolic map representations when needed
- Landscape orientation
- Suitable for serious civil services exam preparation
""".strip()


def generate_and_save_image(article_id: int, topic: str, summary: str):
    t_start = time.time()
    try:
        print(f"[IMAGE] Starting image generation for article {article_id}: {topic}", flush=True)

        prompt = gen_image_prompt(topic, summary)

        t_gen = time.time()
        response = image_client.images.generate(
            model=AZURE_OPENAI_IMAGE_DEPLOYMENT,
            prompt=prompt,
            size="1536x1024"
        )
        gen_duration = round(time.time() - t_gen, 1)

        image_b64 = response.data[0].b64_json
        image_bytes = base64.b64decode(image_b64)

        image = PILImage.open(BytesIO(image_bytes)).convert("RGB")

        image_path = os.path.join(IMAGE_OUTPUT_DIR, f"{article_id}.jpg")

        image.save(
            image_path,
            format="JPEG",
            quality=82,
            optimize=True
        )

        print(f"[IMAGE] Generated image in {gen_duration}s for article {article_id}: {image_path}", flush=True)

        blob_name = f"{article_id}.jpg"
        t_blob = time.time()
        blob_client = blob_service_client.get_blob_client(
            container=AZ_CONTAINER,
            blob=blob_name
        )

        with open(image_path, "rb") as data:
            blob_client.upload_blob(data, overwrite=True)

        blob_url = blob_client.url
        blob_duration = round(time.time() - t_blob, 1)
        total_img_duration = round(time.time() - t_start, 1)

        print(f"[IMAGE] Azure Blob upload completed in {blob_duration}s (Total image pipeline: {total_img_duration}s): {blob_url}", flush=True)

    except Exception as e:
        print(f"[IMAGE] Image generation failed for article {article_id} after {round(time.time() - t_start, 1)}s: {e}", flush=True)


def build_system_prompt() -> str:
    return """
You are an advanced content generator creating markdown material for UPSC aspirants, specifically tailored for database storage and frontend markdown rendering. Your outputs will be referenced by serious exam candidates, so maintain a formal, informative, and concise tone throughout, as if writing official coaching notes or GS exam preparation material. 

Use Bold, underline and Italics as appropriate to highlight important areas (most important don't skip this)

**Task Structure:**
- Use *level-5 markdown headers* (#####) only, never lower or higher.
- Each output consists of strict sections: 
---
##### [Topic Name]
##### [Description in 15-20 words]
<span class="evaluation-result-card-but">GS Paper(s) with topic area e.g. GS III (Economy), GS II (Governance)</span>
(You MUST output exactly ONE span tag like above with relevant UPSC GS papers. Do NOT duplicate or nest span tags.)
##### Key Takeaways  
> [list of highlights from the article important for UPSC examination. Use colored circle emojis like 🔵 🟢 🟠 🔴 🟣 at the start of each bullet point for better cognition]
---
##### Background  
- [Use bullet points to list key facts, historical context, root causes. Each point must start with a hyphen (-) on a new line.]

---
##### Why is it Important for India?
- [Use bullet points to list reasons for national significance, sectors affected, policy impact, impact to institutions. Each point must start with a hyphen (-) on a new line.]

---
##### Data or Statistics  
- [Use bullet points to list up to 5 key statistics. Each point must start with a hyphen (-) on a new line.]

---
##### UPSC - Strategic Viewpoint 
> [Provide 4 or 5 strategic and visionary perspectives like a civil servant explaining the deeper WHY, showing reasoning and judgement, in bullets. Use contextually relevant emojis at the start of each bullet based on the topic - e.g., 🎯 for goals, ⚖️ for policy/law, 🌍 for global, 🇮🇳 for India-specific, 💡 for insights, 📈 for economy, 🛡️ for security, 🌱 for environment, 🏛️ for governance]


---
##### Keywords
[ONLY 5 comma-separated UPSC-relevant keywords including: key terms, constitutional provisions, acts/policies, institutions, concepts, and thematic connections]



- Never add formatting, no block quotes, code blocks, bold/italics, or links in output.

**Content Rules:**
- Synthesize content directly from the provided source article.
- Data or Statistics must be precise, sourced from official or primary documents/sources, and must accurately reflect the correct sequence of events, dates, implementations, and policy mandates for the topic. Always ensure numbers, years, and factual details are consistent.
- Do not repeat output instructions, give meta comments, or include citations/URLs.
- Use proper apostrophes (’) rather than straight (') for all contractions or possessives, ensuring UTF-8 compatibility.
- Be concise, yet insightful and content-rich, distilling only what matters for high-level civil services preparation.
- Write natural, human-like sentences — prioritize clarity, directness, and easy comprehension, avoiding complicated or robotic phrasing.
- STRICTLY DO NOT generate citations, references, footnotes, endnotes, superscripts, or bracketed markers of any kind (e.g. [1], [2], (Source), (Report), etc.)
- The "##### Keywords" section is MANDATORY and must ALWAYS be the last section in your output. You MUST include the exact header "##### Keywords" followed by exactly 5 comma-separated keywords on the next line. Never omit this section — output without it is considered invalid and incomplete.

**Technical:**
- Strictly do NOT wrap output in quotes or code blocks.
- Output must be usable in a SQL or markdown database without post-processing.

**Overall:**
- Prioritize authenticity, clarity, and subject authority in tone and organization.
- Never explain the prompt or process. Output ONLY the formatted markdown content for UPSC use as described.
"""

def is_pib_url(url: str) -> bool:
    parsed = urlparse(url.lower())
    domain = parsed.netloc.replace("www.", "")
    return domain == "pib.gov.in"


def is_bbc_url(url: str) -> bool:
    parsed = urlparse(url.lower())
    domain = parsed.netloc.replace("www.", "")
    return domain == "bbc.com" or domain.endswith("bbc.co.uk")


def is_idsa_url(url: str) -> bool:
    parsed = urlparse(url.lower())
    domain = parsed.netloc.replace("www.", "")
    return domain == "idsa.in"


def is_indianexpress_url(url: str) -> bool:
    parsed = urlparse(url.lower())
    domain = parsed.netloc.replace("www.", "")
    return domain == "indianexpress.com"


def is_thehindu_url(url: str) -> bool:
    parsed = urlparse(url.lower())
    domain = parsed.netloc.replace("www.", "")
    return domain == "thehindu.com"


def is_newindianexpress_url(url: str) -> bool:
    return "newindianexpress.com" in url.lower()


def is_toi_url(url: str) -> bool:
    return "timesofindia.indiatimes.com" in url.lower()


def is_downtoearth_url(url: str) -> bool:
    return "downtoearth.org.in" in url.lower()


def _extract_clean_paragraphs(container, min_len: int = 40):
    paragraphs = []

    for tag in container.find_all([
        'script', 'style', 'aside', 'button', 'nav',
        'header', 'footer', 'figure', 'img'
    ]):
        tag.decompose()

    for p in container.find_all('p'):
        text = p.get_text(" ", strip=True)
        text = re.sub(r'\s+', ' ', text).strip()

        if text and len(text) > min_len and not any(
            noise in text for noise in [
                'Subscribe', 'Advertisement', 'Show Comments',
                'Related Stories', 'Download the', 'Weather', 'Trending'
            ]
        ):
            paragraphs.append(text)

    return paragraphs


def _build_clean_article(soup, article_container, title_selector=None):
    clean_parts = []

    title = None
    if title_selector:
        title = soup.select_one(title_selector)
    if not title:
        title = soup.find('h1')

    if title:
        title_text = title.get_text(" ", strip=True)
        if title_text:
            clean_parts.append(title_text)

    if article_container:
        clean_parts.extend(_extract_clean_paragraphs(article_container))

    cleaned = "\n\n".join(clean_parts).strip()
    return cleaned if len(cleaned) > 150 else ""

def extract_best_content(result, url: str = "") -> str:
    contents = []

    # BBC pages often return heavily hydrated HTML. Extract readable
    # article paragraphs directly instead of returning raw DOM content.
    if url and is_bbc_url(url):
        try:
            html_content = ""
            if hasattr(result, 'cleaned_html') and result.cleaned_html:
                html_content = result.cleaned_html

            if html_content:
                soup = BeautifulSoup(html_content, 'lxml')
                article = soup.find('article')

                if article:
                    paragraphs = []
                    for p in article.find_all('p'):
                        text = p.get_text(" ", strip=True)
                        if text and len(text) > 40:
                            paragraphs.append(text)

                    if paragraphs:
                        return "\n\n".join(paragraphs)
        except Exception:
            pass

    if url and is_idsa_url(url):
        try:
            html_content = ""
            if hasattr(result, 'cleaned_html') and result.cleaned_html:
                html_content = result.cleaned_html

            if html_content:
                soup = BeautifulSoup(html_content, 'lxml')

                clean_parts = []

                title = soup.find('h1', id='posttitle')
                if title:
                    title_text = title.get_text(" ", strip=True)
                    if title_text:
                        clean_parts.append(title_text)

                article_container = (
                    soup.find('div', class_='footnote')
                    or soup.find('div', class_='inner-content-area')
                )

                if article_container:
                    paragraphs = []

                    for p in article_container.find_all('p'):
                        # Remove citation markers like [1], [2]
                        for a in p.find_all('a'):
                            a.decompose()

                        text = p.get_text(" ", strip=True)
                        text = re.sub(r'\[\d+\]', '', text).strip()

                        if text and len(text) > 40:
                            paragraphs.append(text)

                    # Some IDSA articles start with direct text before first <p>
                    direct_text = article_container.get_text(" ", strip=True)
                    if direct_text and len(direct_text) > 100:
                        first_para = direct_text.split("  ")[0].strip()
                        if first_para and first_para not in paragraphs:
                            paragraphs.insert(0, first_para)

                    clean_parts.extend(paragraphs)

                cleaned_content = "\n\n".join(clean_parts).strip()

                if cleaned_content and len(cleaned_content) > 100:
                    return cleaned_content
        except Exception:
            pass

    if url and is_indianexpress_url(url):
        try:
            html_content = ""
            if hasattr(result, 'cleaned_html') and result.cleaned_html:
                html_content = result.cleaned_html

            if html_content:
                soup = BeautifulSoup(html_content, 'lxml')

                clean_parts = []

                title = soup.find('h1', id='main-heading-article')
                subtitle = soup.find('h2', class_='synopsis')

                if title:
                    title_text = title.get_text(" ", strip=True)
                    if title_text:
                        clean_parts.append(title_text)

                if subtitle:
                    subtitle_text = subtitle.get_text(" ", strip=True)
                    if subtitle_text:
                        clean_parts.append(subtitle_text)

                article_container = soup.find('div', id='pcl-full-content')

                if article_container:
                    for tag in article_container.find_all([
                        'script', 'style', 'aside', 'figure', 'img', 'button'
                    ]):
                        tag.decompose()

                    paragraphs = []
                    for p in article_container.find_all('p'):
                        text = p.get_text(" ", strip=True)
                        text = re.sub(r'\s+', ' ', text).strip()

                        if text and len(text) > 40:
                            paragraphs.append(text)

                    clean_parts.extend(paragraphs)

                cleaned_content = "\n\n".join(clean_parts).strip()

                if cleaned_content and len(cleaned_content) > 150:
                    return cleaned_content
        except Exception:
            pass

    if url and is_thehindu_url(url):
        try:
            html_content = ""
            if hasattr(result, 'cleaned_html') and result.cleaned_html:
                html_content = result.cleaned_html

            if html_content:
                soup = BeautifulSoup(html_content, 'lxml')

                clean_parts = []

                title = soup.find('h1')
                if title:
                    title_text = title.get_text(" ", strip=True)
                    if title_text:
                        clean_parts.append(title_text)

                article_container = (
                    soup.select_one('div.articlebodycontent')
                    or soup.select_one('div.storyline')
                    or soup.select_one('div.story-content')
                    or soup.find('article')
                )

                if article_container:
                    for tag in article_container.find_all([
                        'script', 'style', 'aside', 'button', 'nav',
                        'header', 'footer', 'figure', 'img'
                    ]):
                        tag.decompose()

                    paragraphs = []
                    for p in article_container.find_all('p'):
                        text = p.get_text(" ", strip=True)
                        text = re.sub(r'\s+', ' ', text).strip()

                        if (
                            text
                            and len(text) > 40
                            and 'customersupport@thehindu.co.in' not in text
                            and 'Need help with your subscription' not in text
                            and 'Account Settings' not in text
                        ):
                            paragraphs.append(text)

                    clean_parts.extend(paragraphs)

                cleaned_content = "\n\n".join(clean_parts).strip()

                if cleaned_content and len(cleaned_content) > 150:
                    return cleaned_content
        except Exception:
            pass

    if url and is_newindianexpress_url(url):
        try:
            soup = BeautifulSoup(result.cleaned_html or "", 'lxml')
            article = (
                soup.select_one('.full-details')
                or soup.select_one('#pcl-full-content')
                or soup.select_one('.story_details')
            )
            if article:
                cleaned = _build_clean_article(soup, article)
                if cleaned:
                    return cleaned
        except Exception:
            pass

    if url and is_toi_url(url):
        try:
            soup = BeautifulSoup(result.cleaned_html or "", 'lxml')
            article = (
                soup.select_one('div[data-articlebody="1"]')
                or soup.select_one('.article_content')
                or soup.select_one('._s30J')
            )
            if article:
                cleaned = _build_clean_article(soup, article)
                if cleaned:
                    return cleaned
        except Exception:
            pass

    if url and is_downtoearth_url(url):
        try:
            soup = BeautifulSoup(result.cleaned_html or "", 'lxml')
            article = (
                soup.select_one('article')
                or soup.select_one('.story-element-text')
                or soup.select_one('.article-content')
            )
            if article:
                cleaned = _build_clean_article(soup, article)
                if cleaned:
                    return cleaned
        except Exception:
            pass

    try:
        if hasattr(result, 'cleaned_html') and result.cleaned_html:
            soup = BeautifulSoup(result.cleaned_html, 'lxml')

            if url and is_pib_url(url):
                press_div = (
                    soup.find('div', id='divPressRelease')
                    or soup.find('div', id='divPressNote')
                    or soup.find(id='ContentPlaceHolder1_divContent')
                    or soup.select_one('.content-area')
                    or soup.select_one('.main-content')
                    or soup.select_one('.article-content')
                    or soup.select_one('.field-content')
                    or soup.select_one('.content')
                    or soup.find('div', class_='innner-page-main-about-us-content-right-part')
                )

                if press_div:
                    clean_parts = []

                    title = soup.find('h2', id='Titleh2')
                    subtitle = soup.find('h3', id='Subtitleh3')

                    if title:
                        title_text = title.get_text(" ", strip=True)
                        if title_text:
                            clean_parts.append(title_text)

                    if subtitle:
                        subtitle_text = subtitle.get_text(" ", strip=True)
                        if subtitle_text:
                            clean_parts.append(subtitle_text)

                    paragraphs = []
                    for p in press_div.find_all('p'):
                        text = p.get_text(" ", strip=True)
                        if text and len(text) > 30:
                            paragraphs.append(text)

                    clean_parts.extend(paragraphs)

                    cleaned_content = "\n\n".join(clean_parts).strip()

                    if cleaned_content and len(cleaned_content) > 100:
                        return cleaned_content

            common_article_selectors = [
                'article',
                "[role='main']",
                '.article-body',
                '.story-body',
                '.article',
                '.article-content',
                '.post-content',
                '.entry-content',
                '.story-content',
                '.content-body',
                '.description',
                '#DivListing',
                '#content',
                'main',
            ]

            article_container = None
            for selector in common_article_selectors:
                article_container = soup.select_one(selector)
                if article_container:
                    break

            if article_container:
                clean_parts = []

                for tag in article_container.find_all(['script', 'style', 'nav', 'footer', 'header', 'aside']):
                    tag.decompose()

                title_tag = article_container.find(['h1', 'h2'])
                if title_tag:
                    title_text = title_tag.get_text(" ", strip=True)
                    if title_text:
                        clean_parts.append(title_text)

                paragraphs = []
                for p in article_container.find_all('p'):
                    text = p.get_text(" ", strip=True)
                    text = re.sub(r'\s+', ' ', text).strip()
                    text = re.sub(r'\[\d+\]', '', text)

                    if text and len(text) > 40:
                        paragraphs.append(text)

                clean_parts.extend(paragraphs)

                cleaned_content = "\n\n".join(clean_parts).strip()

                if cleaned_content and len(cleaned_content) > 150:
                    return cleaned_content

            tables = soup.find_all('table')
            table_content = "\n".join([t.get_text(separator='\n', strip=True) for t in tables])
            if table_content and len(table_content.strip()) > 100:
                contents.append(("pib_tables", table_content, len(table_content)))
    except Exception:
        pass
    skip_markdown_domains = any([
        is_newindianexpress_url(url),
        is_toi_url(url),
        is_downtoearth_url(url),
    ])

    if not skip_markdown_domains and hasattr(result, 'markdown') and result.markdown:
        if hasattr(result.markdown, 'fit_markdown') and result.markdown.fit_markdown:
            contents.append(("fit_markdown", result.markdown.fit_markdown, len(result.markdown.fit_markdown)))
        if hasattr(result.markdown, 'raw_markdown') and result.markdown.raw_markdown:
            contents.append(("raw_markdown", result.markdown.raw_markdown, len(result.markdown.raw_markdown)))
        if isinstance(result.markdown, str):
            contents.append(("markdown_str", result.markdown, len(result.markdown)))

    if hasattr(result, 'cleaned_html') and result.cleaned_html:
        try:
            soup = BeautifulSoup(result.cleaned_html, 'lxml')

            body = soup.find('body') or soup.find('article') or soup

            for tag in body.find_all([
                'script', 'style', 'nav', 'footer', 'header',
                'aside', 'form', 'button'
            ]):
                tag.decompose()

            extracted_parts = []

            title_tag = body.find(['h1', 'h2', 'title'])
            if title_tag:
                title_text = title_tag.get_text(" ", strip=True)
                if title_text:
                    extracted_parts.append(title_text)

            for p in body.find_all('p'):
                text = p.get_text(" ", strip=True)
                text = re.sub(r'\s+', ' ', text).strip()
                text = re.sub(r'\[\d+\]', '', text)

                if text and len(text) > 30:
                    extracted_parts.append(text)

            cleaned_text = "\n\n".join(extracted_parts).strip()

            if cleaned_text and len(cleaned_text) > 150:
                contents.append(("generic_cleaned_text", cleaned_text, len(cleaned_text)))
            else:
                contents.append(("cleaned_html", result.cleaned_html, len(result.cleaned_html)))

        except Exception:
            contents.append(("cleaned_html", result.cleaned_html, len(result.cleaned_html)))

    contents.sort(key=lambda x: x[2], reverse=True)
    for content_type, content, length in contents:
        if content and length >= 100:
            return content.strip()
    return ""


import httpx


async def fetch_with_fallback(url: str) -> str:
    """Fetch raw page HTML via Wayback Machine archive endpoint when direct fetch fails."""
    async with httpx.AsyncClient(
        timeout=25.0,
        follow_redirects=True,
        headers={
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        }
    ) as client:
        try:
            wayback_api = f"https://archive.org/wayback/available?url={url}"
            api_resp = await client.get(wayback_api)
            if api_resp.status_code == 200:
                data = api_resp.json()
                closest = data.get("archived_snapshots", {}).get("closest")
                if closest and closest.get("available") and closest.get("url"):
                    snapshot_url = closest["url"]
                    page_resp = await client.get(snapshot_url)
                    if page_resp.status_code == 200 and page_resp.text and len(page_resp.text) > 500:
                        return page_resp.text
        except Exception:
            pass

    return ""


async def fetch_article(url: str, prompt: str = "", crawler: AsyncWebCrawler | None = None) -> str:
    browser_config = BrowserConfig(
        headless=True,
        text_mode=False,
        light_mode=True,
        extra_args=[
            "--disable-blink-features=AutomationControlled",
            "--no-first-run",
            "--disable-background-timer-throttling",
            "--disable-renderer-backgrounding",
            "--disable-backgrounding-occluded-windows",
            "--disable-web-security",
            "--disable-features=VizDisplayCompositor",
            "--no-sandbox",
            "--disable-dev-shm-usage"
        ]
    )

    async def _execute_crawl(cfg):
        if crawler is not None:
            return await crawler.arun(url=url, config=cfg)
        else:
            async with AsyncWebCrawler(config=browser_config) as local_crawler:
                return await local_crawler.arun(url=url, config=cfg)

    try:
        if is_pib_url(url):
            crawler_config = CrawlerRunConfig(
                markdown_generator=DefaultMarkdownGenerator(),
                excluded_tags=["nav", "footer", "header", "script", "style"],
                only_text=True,
                cache_mode=CacheMode.BYPASS,
                remove_overlay_elements=True,
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                page_timeout=65000,
                delay_before_return_html=9,
                wait_for="body",
                js_code=[
                    "window.scrollTo(0, document.body.scrollHeight);",
                    "await new Promise(resolve => setTimeout(resolve, 2500));"
                ]
            )
        elif is_bbc_url(url):
            crawler_config = CrawlerRunConfig(
                markdown_generator=DefaultMarkdownGenerator(),
                excluded_tags=[
                    "nav", "footer", "header", "script", "style",
                    "aside", "form", "ads"
                ],
                only_text=True,
                cache_mode=CacheMode.BYPASS,
                remove_overlay_elements=True,
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                page_timeout=75000,
                delay_before_return_html=10,
                wait_for="article",
                css_selector="article",
                js_code=[
                    "window.scrollTo(0, document.body.scrollHeight);",
                    "await new Promise(resolve => setTimeout(resolve, 4000));"
                ]
            )
        elif is_idsa_url(url):
            crawler_config = CrawlerRunConfig(
                markdown_generator=DefaultMarkdownGenerator(),
                excluded_tags=[
                    "nav", "footer", "header", "script", "style",
                    "aside", "form", "ads"
                ],
                only_text=True,
                cache_mode=CacheMode.BYPASS,
                remove_overlay_elements=True,
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                page_timeout=60000,
                delay_before_return_html=9,
                wait_for="body",
                js_code=[
                    "window.scrollTo(0, document.body.scrollHeight);",
                    "await new Promise(resolve => setTimeout(resolve, 7000));"
                ]
            )
        else:
            bm25_filter = BM25ContentFilter(
                user_query=prompt or "current affairs news article",
                bm25_threshold=0.8
            )
            md_generator = DefaultMarkdownGenerator(content_filter=bm25_filter)
            crawler_config = CrawlerRunConfig(
                markdown_generator=md_generator,
                excluded_tags=["nav", "footer", "header", "form", "script", "style", "ads"],
                only_text=True,
                exclude_social_media_links=True,
                keep_data_attributes=False,
                cache_mode=CacheMode.BYPASS,
                remove_overlay_elements=True,
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                page_timeout=75000,
                delay_before_return_html=4,
                wait_for="body"
            )

        result = await _execute_crawl(crawler_config)
        if not result.success:
            raise ValueError(f"Crawling failed: {result.error_message}")
        content = extract_best_content(result, url)
        if content and len(content.strip()) >= 100:
            return content.strip()
        else:
            raise ValueError("Insufficient content extracted")
    except Exception as e:
        try:
            fallback_config = CrawlerRunConfig(
                markdown_generator=DefaultMarkdownGenerator(),
                excluded_tags=["nav", "footer", "header", "script", "style"],
                only_text=True,
                cache_mode=CacheMode.BYPASS,
                remove_overlay_elements=True,
                user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36",
                page_timeout=70000,
                delay_before_return_html=8,
                js_code=[
                    "window.scrollTo(0, document.body.scrollHeight);",
                    "await new Promise(resolve => setTimeout(resolve, 3500));"
                ]
            )
            result = await _execute_crawl(fallback_config)
            content = extract_best_content(result, url)
            if content and len(content.strip()) >= 50:
                return content.strip()
            else:
                ultimate_config = CrawlerRunConfig(
                    markdown_generator=DefaultMarkdownGenerator(),
                    excluded_tags=[],
                    only_text=True,
                    cache_mode=CacheMode.BYPASS,
                    remove_overlay_elements=False,
                    user_agent="Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36",
                    page_timeout=90000,
                    delay_before_return_html=10,
                    js_code=[
                        "window.scrollTo(0, document.body.scrollHeight);",
                        "await new Promise(resolve => setTimeout(resolve, 4500));"
                    ]
                )
                result = await _execute_crawl(ultimate_config)
                content = extract_best_content(result, url)
                if content and len(content.strip()) >= 30:
                    return content.strip()
                else:
                    raise ValueError("No content found in ultimate fallback")
        except Exception as fallback_error:
            if "indianexpress.com" in url or "thehindu.com" in url:
                cached_html = await fetch_with_fallback(url)
                if cached_html:
                    soup = BeautifulSoup(cached_html, "lxml")
                    for tag in soup.find_all(["script", "style", "nav", "footer", "header", "aside", "form"]):
                        tag.decompose()
                    extracted = " ".join(t.get_text(" ", strip=True) for t in soup.find_all(["p", "h1", "h2", "h3"]) if t.get_text(" ", strip=True))
                    extracted = re.sub(r"\s+", " ", extracted).strip()
                    if len(extracted) >= 300 and not _is_blocked_response(extracted):
                        print(f"[FALLBACK] Cache-based extraction succeeded for {url}")
                        return extracted

            raise ValueError(f"All crawling attempts failed. Original: {str(e)}, Ultimate fallback: {str(fallback_error)}")


def generate_content(article_text: str, system_prompt: str) -> str:
    t_llm = time.time()
    client = OpenAI(
        api_key=PPLEX_API_KEY,
        base_url="https://api.perplexity.ai",
        timeout=60.0,
        max_retries=2,
    )
    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": f"""
ARTICLE TO ANALYZE:
{article_text}

Generate comprehensive educational content following the exact template structure provided above. Use your knowledge base to add historical context, relevant data, and background information that enhances the educational value for UPSC aspirants.
"""
        }
    ]
    primary_err = None
    try:
        response = client.chat.completions.create(
            model="sonar-pro",
            messages=messages,
            temperature=0.3,
            max_tokens=3000,
        )
        if response.choices and response.choices[0].message.content:
            duration = round(time.time() - t_llm, 1)
            print(f"[LLM] Perplexity notes generated with sonar-pro in {duration}s", flush=True)
            return response.choices[0].message.content
    except Exception as e1:
        primary_err = e1
        print(f"[LLM] sonar-pro attempt failed ({e1}) after {round(time.time() - t_llm, 1)}s, trying sonar-reasoning...", flush=True)

    try:
        response = client.chat.completions.create(
            model="sonar-reasoning",
            messages=messages,
            temperature=0.3,
            max_tokens=3000,
        )
        if response.choices and response.choices[0].message.content:
            duration = round(time.time() - t_llm, 1)
            print(f"[LLM] Perplexity notes generated with sonar-reasoning in {duration}s", flush=True)
            return response.choices[0].message.content
    except Exception as e2:
        raise ValueError(f"Perplexity API failed on both models. sonar-pro: {primary_err} | sonar-reasoning: {e2}")

    return ""



def clean_and_format_markdown(raw_content: str) -> str:
    """Clean and format the generated markdown content."""
    
    content = re.sub(r'```markdown\n?', '', raw_content)
    content = re.sub(r'\n?```', '', content)
    content = content.replace("'", "'")
    content = re.sub(r'^#{1,6}\s*', '##### ', content, flags=re.MULTILINE)
    content = re.sub(r'https?://\S+', '', content)
    content = re.sub(r'\[(.*?)\]\([^)]*\)', r'\1', content)
    content = re.sub(r'\n\s*\n\s*\n', '\n\n', content)
    return content.strip()

def extract_topic_from_content(content: str) -> str:
    """Extract topic name from the generated content."""
    lines = content.split('\n')
    for line in lines:
        if line.strip().startswith('#####') and line.strip() != '#####':
            topic = line.replace('#####', '').strip()
            if topic and len(topic) < 100:
                return topic
    return "Current_Affairs_Topic"


from typing import Tuple

def split_markdown_sections(md: str) -> Tuple[str, str, str, str]:
    """
    Split markdown into (title, subtitle, remaining_content, keywords).
    Handles optional blank line(s) between title and subtitle.
    Assumes keywords section starts with '##### Keywords'.
    Keywords are returned as a single comma-separated string (as-is).
    """
    lines = md.splitlines()

    # strip leading/trailing blank lines
    while lines and lines[0].strip() == "":
        lines.pop(0)
    while lines and lines[-1].strip() == "":
        lines.pop()

    title, subtitle, remaining_content, keywords = "", "", "", ""

    if not lines:
        return "", "", "", ""

    # always take the first non-empty line as title
    title = lines[0].replace("#####", "").strip()

    # find subtitle (first non-empty line after title)
    subtitle_idx = None
    for i in range(1, len(lines)):
        if lines[i].strip():  # non-empty
            subtitle = lines[i].replace("#####", "").strip()
            subtitle_idx = i
            break

    if subtitle_idx is None:
        # no subtitle, only title
        return title, "", "", ""

    # find the keywords header — robust detection handles LLM variations
    # e.g. "##### Keywords", "##### **Keywords**", "**Keywords**", "Keywords:"
    keyword_idx = None
    for i in range(subtitle_idx + 1, len(lines)):
        line_stripped = re.sub(r'^#+\s*', '', lines[i].strip())   # remove heading markers
        line_stripped = re.sub(r'[\*_]+', '', line_stripped).strip().lower()  # remove bold/italic
        if line_stripped.startswith("keywords"):
            keyword_idx = i
            break

    if keyword_idx is not None:
        remaining_content = "\n".join(lines[subtitle_idx + 1:keyword_idx]).strip()
        # Remove trailing --- separators that precede the Keywords section
        remaining_content = re.sub(r'(\n---\s*)+$', '', remaining_content).strip()
        # Extract keyword text: strip heading markers, bold/italic, and the label itself
        header_text = re.sub(r'^#+\s*', '', lines[keyword_idx])
        header_text = re.sub(r'[\*_]+', '', header_text)
        header_text = re.sub(r'(?i)^keywords\s*:?\s*', '', header_text).strip()
        keyword_lines = []
        if header_text:
            keyword_lines.append(header_text)
        keyword_lines.extend([
            l.strip() for l in lines[keyword_idx+1:]
            if l.strip() and l.strip() != "---"
        ])
        keywords = " ".join(keyword_lines).strip()
    else:
        remaining_content = "\n".join(lines[subtitle_idx + 1:]).strip()

    return title, subtitle, remaining_content, keywords


def preprocess_url(url: str) -> str:
    url = url.strip()

    if "pib.gov.in" in url.lower():
        url = re.sub(r"https?://www\.pib\.gov\.in", "https://pib.gov.in", url, flags=re.IGNORECASE)

        prid_match = re.search(r"PRID=(\d+)", url, flags=re.IGNORECASE)

        if "PressReleaseDetail.aspx" in url and prid_match:
            url = f"https://pib.gov.in/PressReleasePage.aspx?PRID={prid_match.group(1)}"

    if "idsa.in" in url.lower():
        url = re.sub(r"https?://idsa\.in", "https://www.idsa.in", url, flags=re.IGNORECASE)

    if '/amp/' in url or url.endswith('/amp'):
        url = url.replace('/amp/', '/').replace('/amp', '')
    if '?' in url:
        base_url = url.split('?')[0]
        if 'article' in url or 'story' in url:
            return base_url
    return url


def _is_blocked_response(text: str) -> bool:
    if not text:
        return True

    lowered = text.lower()

    block_indicators = [
        "http 403",
        "403 forbidden",
        "forbidden",
        "access denied",
        "access denied",
        "request blocked",
        "blocked by",
        "your ip has been blocked",
        "your request has been blocked",
        "you are not authorized",
        "rate limit",
        "captcha",
        "robot",
        "automated access",
        "cloudflare",
        "waf",
        "access has been restricted",
        "this page could not be loaded",
        "page not found",
        "404 not found",
    ]

    if any(indicator in lowered for indicator in block_indicators):
        return True

    if len(text.strip()) < 300:
        return True

    return False


def classify_topic_subtopic(
    db: orm.Session,
    theme: str,
    title: str,
    subtitle: str,
    keywords: str,
    existing_topic: str | None = None,
) -> tuple[str, str]:
    """
    Classify missing topic/subtopic by querying mtr_subjects, mtr_topics, and mtr_subtopics.
    Uses synthesized article notes (title, subtitle, keywords) for high accuracy and minimal token usage.
    """
    from sqlalchemy import text
    try:
        theme_lower = (theme or "").lower()
        mapping = {
            "ir": ["International Relations", "Internal Security"],
            "security": ["Internal Security", "International Relations"],
            "polity": ["Polity", "Governance"],
            "governance": ["Governance", "Polity"],
            "economy": ["Economy"],
            "geography": ["Geography"],
            "environment": ["Environment/Disaster Management"],
            "disaster": ["Environment/Disaster Management"],
            "science": ["Science & Tech"],
            "tech": ["Science & Tech"],
            "society": ["Society/Social Justice"],
            "social": ["Society/Social Justice"],
            "history": ["Modern History", "Post Independence", "World History"],
            "ethics": ["Ethics"],
            "amac": ["AMAC"],
            "art": ["AMAC"],
            "culture": ["AMAC"],
        }
        matched_subjects = set()
        for k, v in mapping.items():
            if k in theme_lower:
                matched_subjects.update(v)

        if existing_topic and str(existing_topic).strip():
            query = """
                SELECT t.name as topic, GROUP_CONCAT(st.name SEPARATOR ';;;') as subtopics
                FROM mtr_topics t
                LEFT JOIN mtr_subtopics st ON st.topic_id = t.id
                WHERE LOWER(t.name) = LOWER(:topic)
                GROUP BY t.id, t.name
            """
            rows = db.execute(text(query), {"topic": existing_topic.strip()}).fetchall()
            if not rows:
                query = """
                    SELECT t.name as topic, GROUP_CONCAT(st.name SEPARATOR ';;;') as subtopics
                    FROM mtr_topics t
                    LEFT JOIN mtr_subtopics st ON st.topic_id = t.id
                    WHERE t.name LIKE :topic
                    GROUP BY t.id, t.name
                """
                rows = db.execute(text(query), {"topic": f"%{existing_topic.strip()}%"}).fetchall()
        elif matched_subjects:
            placeholders = ",".join(f":s{i}" for i in range(len(matched_subjects)))
            params = {f"s{i}": s for i, s in enumerate(matched_subjects)}
            query = f"""
                SELECT t.name as topic, GROUP_CONCAT(st.name SEPARATOR ';;;') as subtopics
                FROM mtr_subjects s
                JOIN mtr_topics t ON t.subject_id = s.id
                LEFT JOIN mtr_subtopics st ON st.topic_id = t.id
                WHERE s.name IN ({placeholders})
                GROUP BY s.id, t.id, t.name
                ORDER BY s.disp_order, t.disp_order
            """
            rows = db.execute(text(query), params).fetchall()
        else:
            query = """
                SELECT t.name as topic, GROUP_CONCAT(st.name SEPARATOR ';;;') as subtopics
                FROM mtr_subjects s
                JOIN mtr_topics t ON t.subject_id = s.id
                LEFT JOIN mtr_subtopics st ON st.topic_id = t.id
                GROUP BY s.id, t.id, t.name
                ORDER BY s.disp_order, t.disp_order
            """
            rows = db.execute(text(query)).fetchall()

        taxonomy = {}
        for r in rows:
            subs = [st.strip() for st in r[1].split(";;;")] if r[1] else []
            taxonomy[r[0]] = subs

        if not taxonomy:
            return existing_topic or "General", ""

        luna_model = os.getenv("MODEL_GENERATOR", "gpt-5.6-luna")
        azure_chat_client = AzureOpenAI(
            api_key=CA_GEN_AZURE_OPENAI_API_KEY,
            api_version=os.getenv("AZURE_OPENAI_API_VERSION", "2025-01-01-preview"),
            azure_endpoint=CA_GEN_AZURE_OPENAI_ENDPOINT,
            timeout=30.0,
            max_retries=2,
        )

        prompt = f"""Given the following processed UPSC current affairs article:
Article Title: {title}
Summary: {subtitle[:400] if subtitle else ''}
Keywords: {keywords}

Select the SINGLE most appropriate Topic and Subtopic from the following allowed taxonomy list:
{json.dumps(taxonomy, indent=2)}

Return ONLY a JSON object in this exact format:
{{"theme_topic": "<Exact Topic Name>", "theme_subtopic": "<Exact Subtopic Name>"}}
"""

        try:
            resp = azure_chat_client.chat.completions.create(
                model=luna_model,
                messages=[{"role": "user", "content": prompt}],
            )
            raw = resp.choices[0].message.content.strip()
        except Exception as luna_err:
            print(f"[TAXONOMY] gpt-5.6-luna attempt failed ({luna_err}), falling back to sonar-pro...", flush=True)
            pplex_client = OpenAI(
                api_key=PPLEX_API_KEY,
                base_url="https://api.perplexity.ai",
                timeout=25.0,
                max_retries=2,
            )
            resp = pplex_client.chat.completions.create(
                model="sonar-pro",
                messages=[{"role": "user", "content": prompt}],
                temperature=0.1,
            )
            raw = resp.choices[0].message.content.strip()

        match = re.search(r'\{.*?\}', raw, re.DOTALL)
        if match:
            data = json.loads(match.group(0))
            chosen_topic = data.get("theme_topic", "").strip()
            chosen_subtopic = data.get("theme_subtopic", "").strip()

            matched_top = None
            for top in taxonomy:
                if top.lower() == chosen_topic.lower():
                    matched_top = top
                    break
            if not matched_top and existing_topic:
                matched_top = existing_topic

            matched_sub = None
            if matched_top and matched_top in taxonomy:
                allowed_subs = taxonomy[matched_top]
                for sub in allowed_subs:
                    if sub.lower() == chosen_subtopic.lower():
                        matched_sub = sub
                        break
            if not matched_sub:
                matched_sub = chosen_subtopic

            return matched_top or chosen_topic, matched_sub or ""
    except Exception as e:
        print(f"[TAXONOMY] Classification warning: {e}", flush=True)

    return existing_topic or "", ""


async def process_article_service(
    ca_created_date: str | datetime | date,
    theme: str,
    article_url: str,
    db: orm.Session,
    theme_topic: str | None = None,
    theme_subtopic: str | None = None,
    wait_for_image: bool = True,
    crawler: AsyncWebCrawler | None = None,
):

    t_row_start = time.time()

    # 1. Fetch article
    processed_url = preprocess_url(article_url)
    domain = urlparse(processed_url).netloc
    prompt = f"news article from {domain} current affairs"
    t_scrape = time.time()
    article_text = await fetch_article(processed_url, prompt, crawler=crawler)
    scrape_duration = round(time.time() - t_scrape, 1)
    print(f"[SCRAPE] Extracted {len(article_text)} chars in {scrape_duration}s", flush=True)

    if not article_text or len(article_text.strip()) < 50:
        raise ValueError("Insufficient article content retrieved")

    if _is_blocked_response(article_text):
        raise ValueError(f"Source blocked or returned error page (HTTP/403/WAF) for {processed_url}")

    system_prompt = build_system_prompt()
    t_llm = time.time()
    generated_content = generate_content(article_text, system_prompt)
    llm_duration = round(time.time() - t_llm, 1)

    if not generated_content:
        raise ValueError("No content generated by Perplexity")

    clean_content = clean_and_format_markdown(generated_content)

    title, subtitle, rest_content, keywords = split_markdown_sections(clean_content)
    print(f"Keywords: {keywords}", flush=True)

    # Validate extracted title to prevent inserting error pages into the database
    title_check = (title or "").lower()
    invalid_patterns = ["403 forbidden", "http 403", "404 not found", "url not found", "url not", "access denied", "page not found", "blocked by"]
    if any(p in title_check for p in invalid_patterns):
        raise ValueError(f"Generated content is an error page: '{title}'")

    # Classify missing theme_topic or theme_subtopic from mtr_topics/mtr_subtopics
    if not theme_topic or not theme_subtopic or str(theme_topic).strip() == "" or str(theme_subtopic).strip() == "":
        print(f"[TAXONOMY] Classifying missing topic/subtopic against mtr_topics/mtr_subtopics...", flush=True)
        detected_topic, detected_subtopic = classify_topic_subtopic(
            db=db,
            theme=theme,
            title=title,
            subtitle=subtitle,
            keywords=keywords,
            existing_topic=theme_topic,
        )
        if (not theme_topic or str(theme_topic).strip() == "") and detected_topic:
            theme_topic = detected_topic
            print(f"[TAXONOMY] Assigned theme_topic: '{theme_topic}'", flush=True)
        if (not theme_subtopic or str(theme_subtopic).strip() == "") and detected_subtopic:
            theme_subtopic = detected_subtopic
            print(f"[TAXONOMY] Assigned theme_subtopic: '{theme_subtopic}'", flush=True)

    # Parse date strictly from the date specified in the file
    if isinstance(ca_created_date, (datetime, date)):
        parsed_date = datetime(ca_created_date.year, ca_created_date.month, ca_created_date.day)
    elif isinstance(ca_created_date, str) and ca_created_date.strip():
        for fmt in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                parsed_date = datetime.strptime(ca_created_date.strip(), fmt)
                break
            except ValueError:
                pass
        else:
            raise ValueError(f"Could not parse date from file: '{ca_created_date}'")
    else:
        raise ValueError(f"Missing date in file for this article (got {repr(ca_created_date)})")

    # ca_created_date and created_at strictly follow the date mentioned in the file
    created_at = parsed_date.replace(hour=16, minute=7, second=7)

    ca_entry = CurrentAffair(
        topic=title,
        content=subtitle,
        theme=theme,
        theme_topic=theme_topic,
        theme_subtopic=theme_subtopic,
        ca_created_date=parsed_date,
        created_at=created_at,
        source=article_url,
        view_count=random.randint(7, 56)
    )
    db.add(ca_entry)
    db.flush()

    highlight = Highlight(
        id=ca_entry.id,
        current_affair_id=ca_entry.id,
        highlight_text=rest_content,
        created_at=created_at
    )
    db.add(highlight)

    # Build image URL from ca_entry.id
    image_url = f"https://yookicaimages.blob.core.windows.net/thumbnail/{ca_entry.id}.jpg"

    # Store image record
    db.add(Image(
        id=ca_entry.id,
        current_affair_id=ca_entry.id,
        image_url=image_url,
        created_at=created_at
    ))

    # Process keywords
    db.add(Keyword(
        id=ca_entry.id,
        current_affair_id=ca_entry.id,
        kword=keywords,
    ))

    db.commit()
    print(f"[DB] Article record {ca_entry.id} committed successfully", flush=True)

    image_task = None
    if wait_for_image:
        print(f"[IMAGE] Generating and uploading image for article {ca_entry.id}...", flush=True)
        await asyncio.to_thread(
            generate_and_save_image,
            ca_entry.id,
            title,
            subtitle,
        )
    else:
        print(f"[IMAGE] Spawning concurrent background image generation for article {ca_entry.id}", flush=True)
        image_task = asyncio.create_task(
            asyncio.to_thread(
                generate_and_save_image,
                ca_entry.id,
                title,
                subtitle,
            )
        )

    total_article_duration = round(time.time() - t_row_start, 1)
    print(f"[TIMING] Article pipeline completed in {total_article_duration}s (Scrape: {scrape_duration}s | LLM: {llm_duration}s)", flush=True)

    return {
        "id": ca_entry.id,
        "topic": title,
        "theme_topic": theme_topic,
        "theme_subtopic": theme_subtopic,
        "markdown": clean_content,
        "images": [image_url],
        "image_task": image_task,
    }



def classify_failure_reason(exc: Exception, url: str) -> tuple[str, str]:
    """
    Classify failures into human-readable reasons and technical categories.
    Returns (category, reason).
    """
    msg = str(exc)
    msg_lower = msg.lower()

    if not url or not str(url).strip().startswith("http"):
        return "INVALID_LINK", "Link incorrect: Missing or invalid URL format"

    if any(k in msg_lower for k in ["403", "waf", "blocked", "captcha", "cloudflare", "access denied", "forbidden"]):
        return "SCRAPING_BLOCKED", "Scraping not possible: Site blocked access (HTTP 403 / Cloudflare / WAF protection)"

    if "insufficient article content" in msg_lower or "content retrieved" in msg_lower:
        return "CONTENT_INSUFFICIENT", "Scraping not possible: Page body content was empty or insufficient (< 50 chars)"

    if any(k in msg_lower for k in ["perplexity", "sonar", "llm", "no content generated"]):
        return "LLM_SYNTHESIS_FAILED", f"Content generation failed: Perplexity AI error ({msg})"

    if any(k in msg_lower for k in ["image", "openai", "blob", "azure"]):
        return "IMAGE_GENERATION_FAILED", f"Image generation or Azure Blob upload failed ({msg})"

    if any(k in msg_lower for k in ["mysql", "pymysql", "sqlalchemy", "operationalerror", "integrityerror"]):
        return "DATABASE_ERROR", f"Database error: {msg}"

    if any(k in msg_lower for k in ["timeout", "timed out", "connection refused", "name or service not known", "crawling attempts failed"]):
        return "CRAWL_TIMEOUT_OR_NETWORK", f"Scraping not possible: Connection or timeout error ({msg})"

    if "date" in msg_lower:
        return "DATE_ERROR", f"Date error: {msg}"

    return "PROCESSING_ERROR", f"Failed: {msg}"


async def process_excel_file(
    excel_path: str = "ca_lists.xlsx",
    start_row: int = 1,
    limit: int | None = None,
):
    """
    Read an Excel file of current affairs articles, process them sequentially,
    reusing a single browser instance, skipping duplicates, and tracking bulk jobs in mtr_ca_bulk_jobs.
    """
    import openpyxl

    if not os.path.isabs(excel_path):
        excel_path = os.path.join(os.path.dirname(__file__), excel_path)

    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    header_idx = {h.lower().replace(" ", "_").replace("-", "_"): i for i, h in enumerate(headers) if h}

    col_date = header_idx.get("ca_created_date")
    if col_date is None:
        col_date = header_idx.get("date")

    col_theme = header_idx.get("theme")

    col_topic = header_idx.get("theme_topic")
    if col_topic is None:
        col_topic = header_idx.get("topic")

    col_subtopic = header_idx.get("theme_subtopic")
    if col_subtopic is None:
        col_subtopic = header_idx.get("subtopic")

    col_url = header_idx.get("article_url")
    if col_url is None:
        col_url = header_idx.get("url")
    if col_url is None:
        col_url = header_idx.get("link")

    if col_url is None:
        raise ValueError(f"Could not find article_url column in headers: {headers}")

    total_rows = ws.max_row
    print(f"[EXCEL] Loaded {excel_path}. Found {total_rows - 1} rows to inspect.")

    # Initialize status log file
    log_timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_filename = f"ca_status_log_{log_timestamp}.csv"
    log_filepath = os.path.join(os.path.dirname(__file__), log_filename)
    latest_log_path = os.path.join(os.path.dirname(__file__), "ca_status_log_latest.csv")

    fieldnames = [
        "row_num",
        "status",
        "url",
        "date",
        "theme",
        "topic",
        "article_id",
        "failure_category",
        "failure_reason",
        "error_details",
        "processed_at",
    ]

    log_file = open(log_filepath, "w", newline="", encoding="utf-8")
    log_writer = csv.DictWriter(log_file, fieldnames=fieldnames)
    log_writer.writeheader()
    log_file.flush()
    print(f"[LOG] Writing link-by-link status to: {log_filepath}")

    db = get_db_session()

    # Create job in mtr_ca_bulk_jobs
    job_id = str(uuid.uuid4())
    job_started_at = datetime.now()
    bulk_job = BulkJob(
        id=job_id,
        status="processing",
        total=total_rows - 1,
        processed=0,
        succeeded=0,
        failed=0,
        skipped=0,
        details="[]",
        started_at=job_started_at,
        created_at=job_started_at,
    )
    db.add(bulk_job)
    db.commit()
    print(f"[JOB] Registered bulk job {job_id} in mtr_ca_bulk_jobs")

    job_details = []
    processed_count = 0
    success_count = 0
    failed_count = 0
    skipped_count = 0
    pending_image_tasks = []

    browser_config = BrowserConfig(
        headless=True,
        text_mode=False,
        light_mode=True,
        extra_args=[
            "--disable-blink-features=AutomationControlled",
            "--no-first-run",
            "--disable-background-timer-throttling",
            "--disable-renderer-backgrounding",
            "--disable-backgrounding-occluded-windows",
            "--disable-web-security",
            "--disable-features=VizDisplayCompositor",
            "--no-sandbox",
            "--disable-dev-shm-usage"
        ]
    )

    try:
        async with AsyncWebCrawler(config=browser_config) as shared_crawler:
            for row_num in range(2, total_rows + 1):
                if row_num < start_row + 1:
                    continue

                if limit is not None and processed_count >= limit:
                    break

                row_cells = [cell.value for cell in ws[row_num]]

                # Ignore entirely empty trailing rows
                if not any(c is not None and str(c).strip() != "" for c in row_cells):
                    continue

                url_raw = row_cells[col_url] if col_url is not None and col_url < len(row_cells) else None
                date_raw = row_cells[col_date] if col_date is not None and col_date < len(row_cells) else None
                theme_val = str(row_cells[col_theme]).strip() if col_theme is not None and col_theme < len(row_cells) and row_cells[col_theme] else "General"
                topic_val = str(row_cells[col_topic]).strip() if col_topic is not None and col_topic < len(row_cells) and row_cells[col_topic] else None
                subtopic_val = str(row_cells[col_subtopic]).strip() if col_subtopic is not None and col_subtopic < len(row_cells) and row_cells[col_subtopic] else None

                url = str(url_raw).strip() if url_raw is not None else ""

                # Check for link validity
                if not url or not url.startswith("http"):
                    failed_count += 1
                    processed_count += 1
                    cat, reason = "INVALID_LINK", "Link incorrect: URL is missing or not a valid http/https link"
                    print(f"\n[{processed_count}] Row {row_num}: [FAILED] {reason}")
                    job_details.append({
                        "row": row_num,
                        "url": url,
                        "status": "failed",
                        "reason": reason
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "FAILURE",
                        "url": url,
                        "date": str(date_raw) if date_raw else "",
                        "theme": theme_val,
                        "topic": topic_val or "",
                        "article_id": "",
                        "failure_category": cat,
                        "failure_reason": reason,
                        "error_details": f"Raw cell value: {repr(url_raw)}",
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()
                    continue

                # Check for duplicate article already existing in database
                url_clean = preprocess_url(url)
                existing_article = db.query(CurrentAffair).filter(
                    (CurrentAffair.source == url) | (CurrentAffair.source == url_clean)
                ).first()

                if existing_article:
                    skipped_count += 1
                    processed_count += 1
                    dup_reason = f"Duplicate (id={existing_article.id})"
                    print(f"\n[{processed_count}] Row {row_num} ({url}): [SKIPPED] {dup_reason}")
                    job_details.append({
                        "row": row_num,
                        "url": url,
                        "reason": dup_reason,
                        "status": "skipped"
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "SKIPPED",
                        "url": url,
                        "date": str(date_raw) if date_raw else "",
                        "theme": theme_val,
                        "topic": existing_article.topic or topic_val or "",
                        "article_id": existing_article.id,
                        "failure_category": "DUPLICATE",
                        "failure_reason": dup_reason,
                        "error_details": "",
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()
                    continue

                # Check for date presence strictly from file
                if date_raw is None or str(date_raw).strip() == "":
                    failed_count += 1
                    processed_count += 1
                    cat, reason = "MISSING_DATE", f"Date missing in row {row_num}: ca_created_date must be provided in the file"
                    print(f"\n[{processed_count}] Row {row_num} ({url}): [FAILED] {reason}")
                    job_details.append({
                        "row": row_num,
                        "url": url,
                        "status": "failed",
                        "reason": reason
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "FAILURE",
                        "url": url,
                        "date": "",
                        "theme": theme_val,
                        "topic": topic_val or "",
                        "article_id": "",
                        "failure_category": cat,
                        "failure_reason": reason,
                        "error_details": "ca_created_date cell is empty",
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()
                    continue

                # Format date as DD-MM-YYYY string for process_article_service
                if isinstance(date_raw, (datetime, date)):
                    date_val = date_raw.strftime("%d-%m-%Y")
                else:
                    date_val = str(date_raw).strip()

                processed_count += 1
                print(f"\n[{processed_count}] Processing Row {row_num}: {url}")
                print(f"    Date: {date_val} (from file) | Theme: {theme_val} | Topic: {topic_val} | Subtopic: {subtopic_val}")

                try:
                    # Parallel image generation: wait_for_image=False avoids blocking crawler/LLM for next row
                    result = await process_article_service(
                        ca_created_date=date_val,
                        theme=theme_val,
                        article_url=url,
                        db=db,
                        theme_topic=topic_val,
                        theme_subtopic=subtopic_val,
                        wait_for_image=False,
                        crawler=shared_crawler,
                    )
                    if result.get("image_task"):
                        pending_image_tasks.append(result["image_task"])

                    # If theme_topic or theme_subtopic were blank, save classified values back to Excel
                    updated_excel = False
                    resolved_topic = result.get("theme_topic") or topic_val or ""
                    resolved_subtopic = result.get("theme_subtopic") or subtopic_val or ""

                    if col_topic is not None and result.get("theme_topic") and (topic_val is None or str(topic_val).strip() == ""):
                        ws.cell(row=row_num, column=col_topic + 1).value = result["theme_topic"]
                        updated_excel = True
                    if col_subtopic is not None and result.get("theme_subtopic") and (subtopic_val is None or str(subtopic_val).strip() == ""):
                        ws.cell(row=row_num, column=col_subtopic + 1).value = result["theme_subtopic"]
                        updated_excel = True

                    if updated_excel:
                        try:
                            wb.save(excel_path)
                            print(f"[EXCEL] Saved classified topic/subtopic to row {row_num} in {os.path.basename(excel_path)}: Topic='{resolved_topic}' | Subtopic='{resolved_subtopic}'", flush=True)
                        except Exception as save_err:
                            print(f"[EXCEL] Notice: could not save back to Excel: {save_err}", flush=True)

                    success_count += 1
                    print(f"[SUCCESS] Row {row_num} completed -> Article ID: {result['id']} - Topic: {result['topic']} | Theme Topic: {resolved_topic} | Subtopic: {resolved_subtopic}")
                    job_details.append({
                        "id": result["id"],
                        "row": row_num,
                        "url": url,
                        "status": "success",
                        "topic": result.get("topic", ""),
                        "theme_topic": resolved_topic,
                        "theme_subtopic": resolved_subtopic
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "SUCCESS",
                        "url": url,
                        "date": date_val,
                        "theme": theme_val,
                        "topic": resolved_topic or result.get("topic", ""),
                        "article_id": result.get("id", ""),
                        "failure_category": "",
                        "failure_reason": "",
                        "error_details": "",
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()

                except Exception as ex:
                    failed_count += 1
                    db.rollback()
                    cat, reason = classify_failure_reason(ex, url)
                    print(f"[FAILED] Row {row_num} ({url}): [{cat}] {reason}")
                    job_details.append({
                        "row": row_num,
                        "url": url,
                        "status": "failed",
                        "reason": reason
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "FAILURE",
                        "url": url,
                        "date": date_val,
                        "theme": theme_val,
                        "topic": topic_val or "",
                        "article_id": "",
                        "failure_category": cat,
                        "failure_reason": reason,
                        "error_details": str(ex).replace("\n", " ")[:500],
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()

            # Await any concurrent image generation tasks before completing the batch job
            if pending_image_tasks:
                print(f"\n[IMAGE] Awaiting completion of {len(pending_image_tasks)} background image task(s)...", flush=True)
                await asyncio.gather(*pending_image_tasks, return_exceptions=True)
                print(f"[IMAGE] All background image generation and Azure Blob upload tasks completed.", flush=True)


    finally:
        # Update bulk_job in database
        try:
            bulk_job.status = "completed" if failed_count == 0 else "completed_with_errors"
            bulk_job.processed = processed_count
            bulk_job.succeeded = success_count
            bulk_job.failed = failed_count
            bulk_job.skipped = skipped_count
            bulk_job.details = json.dumps(job_details)
            bulk_job.completed_at = datetime.now()
            db.commit()
            print(f"[JOB] Updated job {job_id} in mtr_ca_bulk_jobs")
        except Exception as job_err:
            print(f"[WARN] Failed to update mtr_ca_bulk_jobs: {job_err}")
            db.rollback()

        db.close()
        wb.close()
        log_file.close()
        try:
            shutil.copyfile(log_filepath, latest_log_path)
        except Exception:
            pass

    print(f"\n==================================================")
    print(f"[BATCH COMPLETED]")
    print(f"  Bulk Job ID:          {job_id}")
    print(f"  Total links evaluated: {processed_count}")
    print(f"  Successful:           {success_count}")
    print(f"  Skipped (Duplicate):  {skipped_count}")
    print(f"  Failed:               {failed_count}")
    print(f"  Status Log File:      {log_filepath}")
    print(f"  Latest Log Copy:      {latest_log_path}")
    print(f"==================================================")


if __name__ == "__main__":
    import sys
    target_file = sys.argv[1] if len(sys.argv) > 1 else "ca_lists.xlsx"
    asyncio.run(process_excel_file(target_file))
