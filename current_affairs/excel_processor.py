import asyncio
import csv
import json
import os
import shutil
import uuid
from datetime import date, datetime
import openpyxl

from crawl4ai import AsyncWebCrawler

from current_affairs.config import LOGS_DIR, BASE_DIR
from current_affairs.crawler import preprocess_url, get_default_browser_config
from current_affairs.database import get_db_session, CurrentAffair, BulkJob
from current_affairs.pipeline import process_article_service
from current_affairs.validator import classify_failure_reason


async def process_excel_file(
    excel_path: str = "ca_lists.xlsx",
    start_row: int = 1,
    limit: int | None = None,
) -> str:
    """
    Read an Excel file of current affairs articles, process them sequentially with Crawl4AI,
    run image generation concurrently in the background, classify missing topics via gpt-5.6-luna,
    skip duplicates, and maintain audit logs in logs/ and mtr_ca_bulk_jobs.
    """
    if not os.path.isabs(excel_path):
        excel_path = os.path.join(BASE_DIR, excel_path)

    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    header_idx = {h.lower().replace(" ", "_").replace("-", "_"): i for i, h in enumerate(headers) if h}

    col_date = header_idx.get("ca_created_date")
    if col_date is None:
        col_date = header_idx.get("date")

    col_theme = header_idx.get("theme")

    col_topic = header_idx.get("theme_topic")
    if col_topic is None:
        col_topic = header_idx.get("topic")

    col_subtopic = header_idx.get("theme_subtopic")
    if col_subtopic is None:
        col_subtopic = header_idx.get("subtopic")

    col_url = header_idx.get("article_url")
    if col_url is None:
        col_url = header_idx.get("url")
    if col_url is None:
        col_url = header_idx.get("link")

    if col_url is None:
        raise ValueError(f"Could not find article_url column in headers: {headers}")

    total_rows = ws.max_row
    print(f"[EXCEL] Loaded {excel_path}. Found {total_rows - 1} rows to inspect.", flush=True)

    # Initialize CSV status log inside logs/
    log_timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_filename = f"ca_status_log_{log_timestamp}.csv"
    log_filepath = LOGS_DIR / log_filename
    latest_log_path = LOGS_DIR / "ca_status_log_latest.csv"

    fieldnames = [
        "row_num",
        "status",
        "url",
        "date",
        "theme",
        "topic",
        "article_id",
        "failure_category",
        "failure_reason",
        "error_details",
        "processed_at",
    ]

    log_file = open(log_filepath, "w", newline="", encoding="utf-8")
    log_writer = csv.DictWriter(log_file, fieldnames=fieldnames)
    log_writer.writeheader()
    log_file.flush()
    print(f"[LOG] Writing link-by-link status to: {log_filepath}", flush=True)

    db = get_db_session()

    # Create job in mtr_ca_bulk_jobs
    job_id = str(uuid.uuid4())
    job_started_at = datetime.now()
    bulk_job = BulkJob(
        id=job_id,
        status="processing",
        total=total_rows - 1,
        processed=0,
        succeeded=0,
        failed=0,
        skipped=0,
        details="[]",
        started_at=job_started_at,
        created_at=job_started_at,
    )
    db.add(bulk_job)
    db.commit()
    print(f"[JOB] Registered bulk job {job_id} in mtr_ca_bulk_jobs", flush=True)

    job_details = []
    processed_count = 0
    success_count = 0
    failed_count = 0
    skipped_count = 0
    pending_image_tasks = []

    browser_config = get_default_browser_config()

    try:
        async with AsyncWebCrawler(config=browser_config) as shared_crawler:
            for row_num in range(2, total_rows + 1):
                if row_num < start_row + 1:
                    continue

                if limit is not None and processed_count >= limit:
                    break

                row_cells = [cell.value for cell in ws[row_num]]

                # Ignore entirely empty trailing rows
                if not any(c is not None and str(c).strip() != "" for c in row_cells):
                    continue

                url_raw = row_cells[col_url] if col_url is not None and col_url < len(row_cells) else None
                date_raw = row_cells[col_date] if col_date is not None and col_date < len(row_cells) else None
                theme_val = str(row_cells[col_theme]).strip() if col_theme is not None and col_theme < len(row_cells) and row_cells[col_theme] else "General"
                topic_val = str(row_cells[col_topic]).strip() if col_topic is not None and col_topic < len(row_cells) and row_cells[col_topic] else None
                subtopic_val = str(row_cells[col_subtopic]).strip() if col_subtopic is not None and col_subtopic < len(row_cells) and row_cells[col_subtopic] else None

                url = str(url_raw).strip() if url_raw is not None else ""

                # Check for link validity
                if not url or not url.startswith("http"):
                    failed_count += 1
                    processed_count += 1
                    cat, reason = "INVALID_LINK", "Link incorrect: URL is missing or not a valid http/https link"
                    print(f"\n[{processed_count}] Row {row_num}: [FAILED] {reason}", flush=True)
                    job_details.append({
                        "row": row_num,
                        "url": url,
                        "status": "failed",
                        "reason": reason
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "FAILURE",
                        "url": url,
                        "date": str(date_raw) if date_raw else "",
                        "theme": theme_val,
                        "topic": topic_val or "",
                        "article_id": "",
                        "failure_category": cat,
                        "failure_reason": reason,
                        "error_details": f"Raw cell value: {repr(url_raw)}",
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()
                    continue

                # Check for duplicate article already existing in database
                url_clean = preprocess_url(url)
                existing_article = db.query(CurrentAffair).filter(
                    (CurrentAffair.source == url) | (CurrentAffair.source == url_clean)
                ).first()

                if existing_article:
                    skipped_count += 1
                    processed_count += 1
                    dup_reason = f"Duplicate (id={existing_article.id})"
                    print(f"\n[{processed_count}] Row {row_num} ({url}): [SKIPPED] {dup_reason}", flush=True)
                    job_details.append({
                        "row": row_num,
                        "url": url,
                        "reason": dup_reason,
                        "status": "skipped"
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "SKIPPED",
                        "url": url,
                        "date": str(date_raw) if date_raw else "",
                        "theme": theme_val,
                        "topic": existing_article.topic or topic_val or "",
                        "article_id": existing_article.id,
                        "failure_category": "DUPLICATE",
                        "failure_reason": dup_reason,
                        "error_details": "",
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()
                    continue

                # Check for date presence strictly from file
                if date_raw is None or str(date_raw).strip() == "":
                    failed_count += 1
                    processed_count += 1
                    cat, reason = "MISSING_DATE", f"Date missing in row {row_num}: ca_created_date must be provided in the file"
                    print(f"\n[{processed_count}] Row {row_num} ({url}): [FAILED] {reason}", flush=True)
                    job_details.append({
                        "row": row_num,
                        "url": url,
                        "status": "failed",
                        "reason": reason
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "FAILURE",
                        "url": url,
                        "date": "",
                        "theme": theme_val,
                        "topic": topic_val or "",
                        "article_id": "",
                        "failure_category": cat,
                        "failure_reason": reason,
                        "error_details": "ca_created_date cell is empty",
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()
                    continue

                # Format date as DD-MM-YYYY string for process_article_service
                if isinstance(date_raw, (datetime, date)):
                    date_val = date_raw.strftime("%d-%m-%Y")
                else:
                    date_val = str(date_raw).strip()

                processed_count += 1
                print(f"\n[{processed_count}] Processing Row {row_num}: {url}", flush=True)
                print(f"    Date: {date_val} (from file) | Theme: {theme_val} | Topic: {topic_val} | Subtopic: {subtopic_val}", flush=True)

                try:
                    # Non-blocking image generation: wait_for_image=False avoids blocking crawler/LLM for next row
                    result = await process_article_service(
                        ca_created_date=date_val,
                        theme=theme_val,
                        article_url=url,
                        db=db,
                        theme_topic=topic_val,
                        theme_subtopic=subtopic_val,
                        wait_for_image=False,
                        crawler=shared_crawler,
                    )
                    if result.get("image_task"):
                        pending_image_tasks.append(result["image_task"])

                    # If theme_topic or theme_subtopic were blank, save classified values back to Excel
                    updated_excel = False
                    resolved_topic = result.get("theme_topic") or topic_val or ""
                    resolved_subtopic = result.get("theme_subtopic") or subtopic_val or ""

                    if col_topic is not None and result.get("theme_topic") and (topic_val is None or str(topic_val).strip() == ""):
                        ws.cell(row=row_num, column=col_topic + 1).value = result["theme_topic"]
                        updated_excel = True
                    if col_subtopic is not None and result.get("theme_subtopic") and (subtopic_val is None or str(subtopic_val).strip() == ""):
                        ws.cell(row=row_num, column=col_subtopic + 1).value = result["theme_subtopic"]
                        updated_excel = True

                    if updated_excel:
                        try:
                            wb.save(excel_path)
                            print(f"[EXCEL] Saved classified topic/subtopic to row {row_num} in {os.path.basename(excel_path)}: Topic='{resolved_topic}' | Subtopic='{resolved_subtopic}'", flush=True)
                        except Exception as save_err:
                            print(f"[EXCEL] Notice: could not save back to Excel: {save_err}", flush=True)

                    success_count += 1
                    print(f"[SUCCESS] Row {row_num} completed -> Article ID: {result['id']} - Topic: {result['topic']} | Theme Topic: {resolved_topic} | Subtopic: {resolved_subtopic}", flush=True)
                    job_details.append({
                        "id": result["id"],
                        "row": row_num,
                        "url": url,
                        "status": "success",
                        "topic": result.get("topic", ""),
                        "theme_topic": resolved_topic,
                        "theme_subtopic": resolved_subtopic,
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "SUCCESS",
                        "url": url,
                        "date": date_val,
                        "theme": theme_val,
                        "topic": resolved_topic or result.get("topic", ""),
                        "article_id": result.get("id", ""),
                        "failure_category": "",
                        "failure_reason": "",
                        "error_details": "",
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()

                except Exception as ex:
                    failed_count += 1
                    db.rollback()
                    cat, reason = classify_failure_reason(ex, url)
                    print(f"[FAILED] Row {row_num} ({url}): [{cat}] {reason}", flush=True)
                    job_details.append({
                        "row": row_num,
                        "url": url,
                        "status": "failed",
                        "reason": reason
                    })
                    log_writer.writerow({
                        "row_num": row_num,
                        "status": "FAILURE",
                        "url": url,
                        "date": date_val,
                        "theme": theme_val,
                        "topic": topic_val or "",
                        "article_id": "",
                        "failure_category": cat,
                        "failure_reason": reason,
                        "error_details": str(ex).replace("\n", " ")[:500],
                        "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })
                    log_file.flush()

            # Await any concurrent image generation tasks before completing the batch job
            if pending_image_tasks:
                print(f"\n[IMAGE] Awaiting completion of {len(pending_image_tasks)} background image task(s)...", flush=True)
                await asyncio.gather(*pending_image_tasks, return_exceptions=True)
                print(f"[IMAGE] All background image generation and Azure Blob upload tasks completed.", flush=True)

    finally:
        # Update bulk_job in database
        try:
            bulk_job.status = "completed" if failed_count == 0 else "completed_with_errors"
            bulk_job.processed = processed_count
            bulk_job.succeeded = success_count
            bulk_job.failed = failed_count
            bulk_job.skipped = skipped_count
            bulk_job.details = json.dumps(job_details)
            bulk_job.completed_at = datetime.now()
            db.commit()
            print(f"[JOB] Updated job {job_id} in mtr_ca_bulk_jobs", flush=True)
        except Exception as job_err:
            print(f"[WARN] Failed to update mtr_ca_bulk_jobs: {job_err}", flush=True)
            db.rollback()

        db.close()
        wb.close()
        log_file.close()
        try:
            shutil.copyfile(str(log_filepath), str(latest_log_path))
        except Exception:
            pass

    print(f"\n==================================================")
    print(f"[BATCH COMPLETED]")
    print(f"  Bulk Job ID:          {job_id}")
    print(f"  Total links evaluated: {processed_count}")
    print(f"  Successful:           {success_count}")
    print(f"  Skipped (Duplicate):  {skipped_count}")
    print(f"  Failed:               {failed_count}")
    print(f"  Status Log File:      {log_filepath}")
    print(f"  Latest Log Copy:      {latest_log_path}")
    print(f"==================================================")
    return job_id
